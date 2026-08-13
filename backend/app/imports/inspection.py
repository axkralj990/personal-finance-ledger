from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import struct
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from openpyxl.utils.exceptions import InvalidFileException

from backend.app.imports.models import (
    ImportFileType,
    ImportInspection,
    InferredValueType,
    InspectedColumn,
    SheetInspection,
    SourceRow,
)

MAX_SHEETS = 32
MAX_COLUMNS = 256
MAX_INSPECTED_CELLS = 5_000_000
MAX_UNCOMPRESSED_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_PREVIEW_ROWS = 100
MAX_HEADER_SCAN_ROWS = 50
MAX_CELL_CHARACTERS = 32_767
XLS_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
UNSAFE_XLSX_PATHS = (
    re.compile(r"(^|/)vbaProject\.bin$", re.IGNORECASE),
    re.compile(r"^xl/externalLinks/", re.IGNORECASE),
    re.compile(r"^xl/connections\.xml$", re.IGNORECASE),
    re.compile(r"^xl/queryTables/", re.IGNORECASE),
)


class InspectionError(ValueError):
    pass


def inspect_file(
    path: Path,
    *,
    sheet: str | int | None = None,
    header_row: int | None = None,
    preview_offset: int = 0,
    preview_limit: int = 100,
    max_rows: int = 100_000,
) -> ImportInspection:
    if preview_offset < 0:
        raise InspectionError("preview_offset cannot be negative")
    if not 1 <= preview_limit <= MAX_PREVIEW_ROWS:
        raise InspectionError(f"preview_limit must be between 1 and {MAX_PREVIEW_ROWS}")
    if max_rows < 1:
        raise InspectionError("max_rows must be positive")

    file_type = detect_file_type(path)
    if file_type == ImportFileType.CSV:
        tables, encoding, delimiter, quote_character = _read_csv(path, max_rows)
    elif file_type == ImportFileType.XLSX:
        tables = _read_xlsx(path, max_rows)
        encoding = delimiter = quote_character = None
    else:
        tables = _read_xls(path, max_rows)
        encoding = delimiter = quote_character = None

    summaries = tuple(
        _sheet_summary(index, name, rows) for index, (name, rows) in enumerate(tables)
    )
    selected_index = _select_sheet(tables, sheet)
    selected_name, rows = tables[selected_index]
    selected_header = _select_header(rows, header_row)
    width = _table_width(rows, selected_header)
    if width == 0:
        raise InspectionError("selected table has no columns")
    if width > MAX_COLUMNS:
        raise InspectionError(f"selected table exceeds {MAX_COLUMNS} columns")

    padded_header = _pad_row(rows[selected_header - 1], width)
    data_rows = [
        (row_number, _pad_row(row, width))
        for row_number, row in enumerate(rows[selected_header:], start=selected_header + 1)
    ]
    non_trailing = [index for index, (_, row) in enumerate(data_rows) if not _is_blank_row(row)]
    data_end_row = data_rows[non_trailing[-1]][0] if non_trailing else selected_header
    bounded_data = [(number, row) for number, row in data_rows if number <= data_end_row]
    if len(bounded_data) > max_rows:
        raise InspectionError(f"selected table exceeds {max_rows} data rows")

    blank_rows = tuple(number for number, row in bounded_data if _is_blank_row(row))
    repeated_headers = tuple(
        number
        for number, row in bounded_data
        if _normalized_row(row) == _normalized_row(padded_header)
    )
    possible_footers = tuple(
        number
        for number, row in bounded_data[-10:]
        if number not in blank_rows and number not in repeated_headers and _looks_like_footer(row)
    )
    sample_rows = [row for number, row in bounded_data if number not in blank_rows][:200]
    columns = tuple(
        InspectedColumn(
            id=f"c{position:03d}",
            position=position,
            raw_label=_cell_text(label),
            normalized_label=normalize_label(label),
            inferred_type=_infer_type([row[position] for row in sample_rows]),
        )
        for position, label in enumerate(padded_header)
    )
    preview_rows = bounded_data[preview_offset : preview_offset + preview_limit]
    preview = tuple(_source_row(number, row) for number, row in preview_rows)
    signature = structural_signature(
        file_type=file_type,
        selected_sheet=selected_name,
        sheet_index=selected_index,
        header_row=selected_header,
        columns=columns,
    )
    return ImportInspection(
        file_type=file_type,
        encoding=encoding,
        delimiter=delimiter,
        quote_character=quote_character,
        sheets=summaries,
        selected_sheet=selected_name,
        header_row=selected_header,
        data_start_row=selected_header + 1,
        data_end_row=data_end_row,
        row_count=len(bounded_data),
        blank_rows=blank_rows,
        repeated_header_rows=repeated_headers,
        possible_footer_rows=possible_footers,
        columns=columns,
        preview_offset=preview_offset,
        preview=preview,
        structural_signature=signature,
    )


def read_source_rows(
    path: Path,
    inspection: ImportInspection,
    *,
    max_rows: int = 100_000,
) -> list[SourceRow]:
    current = inspect_file(
        path,
        sheet=inspection.selected_sheet,
        header_row=inspection.header_row,
        preview_limit=1,
        max_rows=max_rows,
    )
    if current.structural_signature != inspection.structural_signature:
        raise InspectionError("file structure no longer matches the inspection")
    file_type = current.file_type
    if file_type == ImportFileType.CSV:
        tables, _, _, _ = _read_csv(path, max_rows)
    elif file_type == ImportFileType.XLSX:
        tables = _read_xlsx(path, max_rows)
    else:
        tables = _read_xls(path, max_rows)
    rows = dict(tables)[inspection.selected_sheet]
    width = len(inspection.columns)
    return [
        _source_row(number, _pad_row(row, width))
        for number, row in enumerate(rows[inspection.header_row :], inspection.data_start_row)
        if number <= inspection.data_end_row
    ]


def detect_file_type(path: Path) -> ImportFileType:
    try:
        prefix = path.read_bytes()[:8]
    except OSError as exc:
        raise InspectionError("file could not be read") from exc
    if prefix.startswith(XLS_SIGNATURE):
        return ImportFileType.XLS
    if prefix.startswith(ZIP_SIGNATURES):
        _validate_xlsx_archive(path)
        return ImportFileType.XLSX
    sample = path.read_bytes()[:4096]
    if b"\x00" in sample and not sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise InspectionError("unsupported or malformed binary file")
    return ImportFileType.CSV


def structural_signature(
    *,
    file_type: ImportFileType,
    selected_sheet: str,
    sheet_index: int,
    header_row: int,
    columns: Sequence[InspectedColumn],
) -> str:
    payload = {
        "inspection_version": "inspection-v1",
        "execution_schema_version": "universal-v1",
        "file_type": file_type.value,
        "sheet": {"index": sheet_index, "name": selected_sheet},
        "header_row": header_row,
        "columns": [
            {"id": column.id, "position": column.position, "label": column.normalized_label}
            for column in columns
        ],
    }
    encoded = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode()).hexdigest()


def normalize_label(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _cell_text(value)).casefold().strip()
    return " ".join(re.sub(r"[^\w]+", " ", text, flags=re.UNICODE).split())


def _read_csv(path: Path, max_rows: int) -> tuple[list[tuple[str, list[list[Any]]]], str, str, str]:
    raw = path.read_bytes()
    if not raw:
        raise InspectionError("file is empty")
    text, encoding = _decode_csv(raw)
    sample = text[:64_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = []
        reader = csv.reader(io.StringIO(text, newline=""), dialect, strict=True)
        for row in reader:
            if len(row) > MAX_COLUMNS:
                raise InspectionError(f"CSV exceeds {MAX_COLUMNS} columns")
            rows.append([_bounded_cell(value) for value in row])
            if len(rows) > max_rows + MAX_HEADER_SCAN_ROWS + 1:
                raise InspectionError(f"CSV exceeds {max_rows} data rows")
    except csv.Error as exc:
        raise InspectionError("CSV is malformed") from exc
    if not rows:
        raise InspectionError("CSV has no rows")
    return [("CSV", rows)], encoding, dialect.delimiter, dialect.quotechar


def _decode_csv(raw: bytes) -> tuple[str, str]:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        candidates = ("utf-16",)
    else:
        candidates = ("utf-8-sig", "cp1252")
    for encoding in candidates:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if text and text.count("\ufffd") == 0:
            return text, encoding
    raise InspectionError("CSV encoding is unsupported or malformed")


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 20_000:
                raise InspectionError("workbook archive contains too many entries")
            total = 0
            names = set()
            for info in infos:
                name = info.filename.replace("\\", "/")
                if name.startswith("/") or ".." in name.split("/"):
                    raise InspectionError("workbook archive contains an unsafe path")
                total += info.file_size
                if total > MAX_UNCOMPRESSED_WORKBOOK_BYTES:
                    raise InspectionError("workbook uncompressed content exceeds 100 MiB")
                if info.compress_size == 0 and info.file_size > 1_000_000:
                    raise InspectionError("workbook archive has an unsafe compression ratio")
                if info.compress_size and info.file_size / info.compress_size > 1_000:
                    raise InspectionError("workbook archive has an unsafe compression ratio")
                names.add(name)
                if any(pattern.search(name) for pattern in UNSAFE_XLSX_PATHS):
                    raise InspectionError(
                        "workbook contains macros, external links, or connections"
                    )
                if info.flag_bits & 0x1:
                    raise InspectionError("encrypted workbooks are not supported")
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset(names):
                raise InspectionError("ZIP file is not an XLSX workbook")
            content_types = archive.read("[Content_Types].xml")
            if b"macroEnabled" in content_types or b"vbaProject" in content_types:
                raise InspectionError("workbook contains macros")
            bad_member = archive.testzip()
            if bad_member is not None:
                raise InspectionError("workbook archive is malformed")
    except (zipfile.BadZipFile, OSError) as exc:
        raise InspectionError("workbook archive is malformed") from exc


def _read_xlsx(path: Path, max_rows: int) -> list[tuple[str, list[list[Any]]]]:
    _validate_xlsx_archive(path)
    stream = path.open("rb")
    try:
        workbook = openpyxl.load_workbook(
            stream,
            read_only=True,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        stream.close()
        raise InspectionError("XLSX workbook is malformed or encrypted") from exc
    try:
        return _extract_xlsx_tables(workbook, max_rows)
    except InspectionError:
        raise
    except Exception as exc:
        raise InspectionError("XLSX workbook is malformed") from exc
    finally:
        workbook.close()
        stream.close()


def _extract_xlsx_tables(
    workbook: openpyxl.Workbook, max_rows: int
) -> list[tuple[str, list[list[Any]]]]:
    if len(workbook.sheetnames) > MAX_SHEETS:
        raise InspectionError(f"workbook exceeds {MAX_SHEETS} sheets")
    tables = []
    total_cells = 0
    for worksheet in workbook.worksheets:
        rows: list[list[Any]] = []
        for row_index, cells in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = [_bounded_cell(value) for value in cells]
            while values and values[-1] is None:
                values.pop()
            if len(values) > MAX_COLUMNS:
                raise InspectionError(f"sheet {worksheet.title!r} exceeds {MAX_COLUMNS} columns")
            total_cells += len(values)
            if total_cells > MAX_INSPECTED_CELLS:
                raise InspectionError("workbook exceeds 5,000,000 inspected cells")
            rows.append(values)
            if row_index > max_rows + MAX_HEADER_SCAN_ROWS + 1:
                raise InspectionError(f"sheet {worksheet.title!r} exceeds {max_rows} data rows")
        tables.append((worksheet.title, rows))
    return tables


def _read_xls(path: Path, max_rows: int) -> list[tuple[str, list[list[Any]]]]:
    raw = path.read_bytes()
    if len(raw) > MAX_UNCOMPRESSED_WORKBOOK_BYTES:
        raise InspectionError("workbook content exceeds 100 MiB")
    try:
        compound_document = xlrd.compdoc.CompDoc(raw, logfile=io.StringIO())
    except (CompDocError, IndexError, struct.error, ValueError) as exc:
        raise InspectionError("XLS workbook is malformed or encrypted") from exc
    stream_names = {entry.name.casefold() for entry in compound_document.dirlist}
    if any("vba" in name or "macros" in name for name in stream_names):
        raise InspectionError("workbook contains macros")
    if stream_names & {"encryptedpackage", "encryptioninfo"}:
        raise InspectionError("encrypted workbooks are not supported")
    try:
        workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
    except (xlrd.XLRDError, CompDocError, OSError, struct.error, ValueError) as exc:
        raise InspectionError("XLS workbook is malformed or encrypted") from exc
    try:
        return _extract_xls_tables(workbook, max_rows)
    except InspectionError:
        raise
    except Exception as exc:
        raise InspectionError("XLS workbook is malformed") from exc
    finally:
        workbook.release_resources()


def _extract_xls_tables(
    workbook: xlrd.book.Book, max_rows: int
) -> list[tuple[str, list[list[Any]]]]:
    if workbook.nsheets > MAX_SHEETS:
        raise InspectionError(f"workbook exceeds {MAX_SHEETS} sheets")
    unsafe_supbooks = {xlrd.book.SUPBOOK_EXTERNAL, xlrd.book.SUPBOOK_DDEOLE}
    if any(kind in unsafe_supbooks for kind in workbook._supbook_types):
        raise InspectionError("workbook contains external links or connections")
    tables = []
    total_cells = 0
    for sheet in workbook.sheets():
        if sheet.ncols > MAX_COLUMNS:
            raise InspectionError(f"sheet {sheet.name!r} exceeds {MAX_COLUMNS} columns")
        if sheet.nrows > max_rows + MAX_HEADER_SCAN_ROWS + 1:
            raise InspectionError(f"sheet {sheet.name!r} exceeds {max_rows} data rows")
        total_cells += sheet.nrows * sheet.ncols
        if total_cells > MAX_INSPECTED_CELLS:
            raise InspectionError("workbook exceeds 5,000,000 inspected cells")
        rows = [
            [
                _bounded_cell(_xls_value(sheet.cell(row, column), workbook.datemode))
                for column in range(sheet.ncols)
            ]
            for row in range(sheet.nrows)
        ]
        tables.append((sheet.name, rows))
    return tables


try:
    from xlrd.compdoc import CompDocError
except ImportError:  # pragma: no cover - supported xlrd exposes this error.
    CompDocError = xlrd.XLRDError


def _xls_value(cell: xlrd.sheet.Cell, datemode: int) -> Any:
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return None
    return cell.value


def _sheet_summary(index: int, name: str, rows: list[list[Any]]) -> SheetInspection:
    width = max((len(row) for row in rows), default=0)
    return SheetInspection(
        name=name,
        index=index,
        row_count=len(rows),
        column_count=min(width, MAX_COLUMNS),
        candidate_header_rows=tuple(_candidate_headers(rows)),
    )


def _candidate_headers(rows: list[list[Any]]) -> list[int]:
    scored = []
    for index, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS], start=1):
        non_empty = [value for value in row if not _is_blank(value)]
        if len(non_empty) < 2:
            continue
        text_count = sum(
            not _looks_number(value) and not isinstance(value, date) for value in non_empty
        )
        unique = len({_cell_text(value).casefold() for value in non_empty})
        next_non_empty = sum(
            not _is_blank(value)
            for candidate in rows[index : index + 5]
            for value in candidate[: len(row)]
        )
        score = len(non_empty) * 4 + text_count * 3 + unique + min(next_non_empty, 20)
        scored.append((score, -index, index))
    return [item[2] for item in sorted(scored, reverse=True)[:3]]


def _select_sheet(tables: list[tuple[str, list[list[Any]]]], choice: str | int | None) -> int:
    if not tables:
        raise InspectionError("workbook has no sheets")
    if isinstance(choice, int):
        if not 0 <= choice < len(tables):
            raise InspectionError("selected sheet index does not exist")
        return choice
    if isinstance(choice, str):
        for index, (name, _) in enumerate(tables):
            if name == choice:
                return index
        raise InspectionError("selected sheet does not exist")
    candidates = [
        (len(_candidate_headers(rows)) > 0, sum(len(row) for row in rows), -index, index)
        for index, (_, rows) in enumerate(tables)
    ]
    return max(candidates)[3]


def _select_header(rows: list[list[Any]], choice: int | None) -> int:
    if choice is not None:
        if not 1 <= choice <= len(rows):
            raise InspectionError("selected header row does not exist")
        return choice
    candidates = _candidate_headers(rows)
    if not candidates:
        raise InspectionError("no candidate header row was detected")
    return candidates[0]


def _table_width(rows: list[list[Any]], header_row: int) -> int:
    relevant = rows[header_row - 1 :]
    return max((len(row) for row in relevant), default=0)


def _pad_row(row: Sequence[Any], width: int) -> list[Any]:
    return [*row[:width], *([None] * max(0, width - len(row)))]


def _source_row(row_number: int, row: Sequence[Any]) -> SourceRow:
    return SourceRow(
        row_number=row_number,
        values={f"c{position:03d}": _json_value(value) for position, value in enumerate(row)},
    )


def _bounded_cell(value: Any) -> Any:
    if isinstance(value, str) and len(value) > MAX_CELL_CHARACTERS:
        raise InspectionError(f"cell text exceeds {MAX_CELL_CHARACTERS} characters")
    return value


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, str | int | bool):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) else value
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return not _cell_text(value)


def _is_blank_row(row: Sequence[Any]) -> bool:
    return all(_is_blank(value) for value in row)


def _normalized_row(row: Sequence[Any]) -> tuple[str, ...]:
    return tuple(normalize_label(value) for value in row)


def _looks_like_footer(row: Sequence[Any]) -> bool:
    values = [normalize_label(value) for value in row if not _is_blank(value)]
    if not values:
        return False
    markers = ("total", "subtotal", "balance", "opening balance", "closing balance", "summary")
    return len(values) <= 3 and any(value.startswith(markers) for value in values)


def _looks_number(value: Any) -> bool:
    if isinstance(value, int | float) and not isinstance(value, bool):
        return True
    text = _cell_text(value)
    return bool(
        re.fullmatch(
            r"[()\-+]?\s*[$€£¥]?\s*\d[\d\s.,']*\s*[$€£¥]?\s*-?",
            text,
        )
    )


def _infer_type(values: Iterable[Any]) -> InferredValueType:
    non_empty = [value for value in values if not _is_blank(value)][:100]
    if not non_empty:
        return InferredValueType.EMPTY
    kinds = {_value_kind(value) for value in non_empty}
    if len(kinds) == 1:
        return kinds.pop()
    if kinds <= {InferredValueType.NUMBER, InferredValueType.CURRENCY}:
        return InferredValueType.CURRENCY
    return InferredValueType.MIXED


def _value_kind(value: Any) -> InferredValueType:  # noqa: PLR0911
    if isinstance(value, datetime):
        return InferredValueType.TIMESTAMP
    if isinstance(value, date):
        return InferredValueType.DATE
    text = _cell_text(value)
    if _looks_number(value):
        if any(symbol in text for symbol in "$€£¥"):
            return InferredValueType.CURRENCY
        return InferredValueType.NUMBER
    iso_candidate = text.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(iso_candidate)
    except ValueError:
        return InferredValueType.TEXT
    if parsed.time().isoformat() != "00:00:00":
        return InferredValueType.TIMESTAMP
    return InferredValueType.DATE
